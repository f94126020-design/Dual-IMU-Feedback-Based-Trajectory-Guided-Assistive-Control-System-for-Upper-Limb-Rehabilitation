from docx import Document
import json

path = r"C:\Users\USER\Desktop\NCKU\推甄準備\成大醫工醫學電子組\成大醫工研究計畫_最終版.docx"
doc = Document(path)

for i, p in enumerate(doc.paragraphs):
    if any(k in p.text for k in ("IC", "晶片", "積體電路")):
        print(json.dumps({"kind": "paragraph", "index": i, "text": p.text}, ensure_ascii=True))

for ti, table in enumerate(doc.tables):
    for ri, row in enumerate(table.rows):
        for ci, cell in enumerate(row.cells):
            if any(k in cell.text for k in ("IC", "晶片", "積體電路")):
                print(json.dumps({"kind": "cell", "table": ti, "row": ri, "col": ci, "text": cell.text}, ensure_ascii=True))
