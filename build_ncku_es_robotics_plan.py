from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from format_research_plans import format_document


SRC = Path(r"C:\Users\USER\Desktop\NCKU\推甄準備\研究計畫\A版_成大電機控制組_穿戴式IMU可信度感知遙操作研究計畫.docx")
OUT = Path(r"C:\Users\USER\Desktop\NCKU\推甄準備\研究計畫\A版_成大工科系統整合暨智慧機器人組_研究計畫.docx")


def replace_paragraph(paragraph, text):
    paragraph.clear()
    paragraph.add_run(text)


def remove_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)


doc = Document(SRC)

replace_paragraph(doc.paragraphs[0], "穿戴式 IMU 可信度感知之機器人遙操作與嵌入式安全控制系統")
replace_paragraph(doc.paragraphs[1], "成大工科系統整合暨智慧機器人組版　感測、控制、通訊與機電系統整合")

replace_paragraph(doc.paragraphs[4], "機器人透過穿戴式介面接受人類遙操作命令時，感測、通訊與控制模組必須共同維持即時性與安全性。若慣性量測單元（IMU）受到零偏、累積漂移、穿戴位置改變、機構振動或無線封包遺失影響，錯誤量測可能沿著資料鏈轉換成非預期機器人運動。現有研究已證明 IMU 遙操作的可行性 [1–4]，並分別處理多 IMU 校正 [5]、穿戴校正 [8–9] 與不確定性感知共享控制 [6–7]；然而，較少從完整機器人系統角度，將感測可信度、無線通訊狀態、安全控制律與嵌入式故障管理整合於同一平台並進行實體驗證。本研究將建立穿戴式 IMU 可信度估測器，依即時可靠度調整操作者命令權重、速度與安全限制，並在 MCU、嵌入式主機及機械手臂平台上，量化感測異常、通訊延遲與系統負載對任務表現及安全行為的影響。")

replace_paragraph(doc.paragraphs[6], "智慧機器人並非單一控制器或辨識模型，而是由感測器、通訊、嵌入式運算、控制器、致動器及安全機制共同構成的機電系統。一般遙操作架構常將穿戴式量測姿態直接映射為機器人參考命令；一旦裝置鬆動、重新穿戴、感測器漂移或無線傳輸中斷，即使各模組個別運作正常，整體系統仍可能追蹤錯誤命令。因此，本研究關注的不是單獨提高 IMU 準確度，而是機器人系統如何辨識輸入資料的可靠程度，並協調感測、控制與故障管理，使機器人在非理想條件下仍能安全退讓。")

replace_paragraph(doc.paragraphs[7], "這項問題源自我的畢業專題經驗。我目前就讀生物醫學工程學系四年級，在雙 IMU 上肢輔助系統中參與姿態量測、Raspberry Pi 軟體、ESP32 韌體、馬達控制、通訊與安全狀態機整合。實作時，我觀察到馬達與機構振動會回饋至 IMU 訊號，也體會到取樣、封包、控制週期與致動安全彼此牽連。這段經驗使我希望進一步研究跨感測、控制與嵌入式平台的智慧機器人系統，並累積機器人核心技術、機電整合、人機協作與智慧物聯網所需的系統實作能力。")

for p in doc.paragraphs:
    if p.text == "三、研究問題與概念模型":
        replace_paragraph(p, "三、研究問題與系統架構")
    elif p.text.startswith("如何設計以 IMU 可信度為調度依據的共享控制律"):
        replace_paragraph(p, "如何以 IMU 即時可信度協調人類命令、機器人控制權與安全限制，使系統在感測或通訊異常下抑制非預期運動？")
    elif p.text == "4.3 人機姿態映射與安全約束":
        replace_paragraph(p, "4.3 人機命令映射與機器人安全約束")
    elif p.text == "4.4 可信度相依共享控制":
        replace_paragraph(p, "4.4 可信度感知控制與系統協調")
    elif p.text.startswith("本研究以可信度作為控制器的調度訊號"):
        replace_paragraph(p, "本研究以可信度作為系統協調訊號，連續調整操作者命令權重與允許速度，並加入關節、工作空間、輸入飽和及命令斜率限制。控制器比較固定增益與可信度感知策略，分析感測異常、通訊延遲與狀態切換對軌跡誤差、輸出平順性及安全邊界的影響。系統設計同時定義正常、降權、保持與停止等運作模式，以及各模式間的進入、恢復與逾時條件，使估測結果能轉換為可測試、可追蹤的機器人行為，而不只停留在異常偵測。")
    elif p.text == "4.5 嵌入式實現與故障管理":
        replace_paragraph(p, "4.5 嵌入式韌體、通訊與故障管理")
    elif p.text.startswith("將感測、姿態估測、可信度更新與命令輸出分成固定週期模組"):
        replace_paragraph(p, "系統以 ESP32 或同級 MCU 負責固定週期取樣、時間戳、封包序號、watchdog、timeout 與急停，以 Raspberry Pi 或嵌入式主機執行姿態估測、可信度更新、控制與資料記錄，並透過 ROS 或模組化介面連接機器人平台。各模組將定義資料格式、更新率、有效性旗標及錯誤碼，記錄端到端延遲、最壞執行時間與控制迴圈抖動。故障測試涵蓋資料過期、封包遺失、緩衝區溢位、節點重啟及低可信度持續狀態，以驗證跨裝置系統能否安全降級並恢復。")
    elif p.text == "九、申請成大電機控制組之研究契合":
        replace_paragraph(p, "九、與成大工科系統整合暨智慧機器人組之契合")
    elif p.text.startswith("成大電機控制組的研究包含"):
        replace_paragraph(p, "成大工程科學系重視理論基礎、系統思維與跨領域實作，機器人相關訓練涵蓋感測、自動控制、程式設計、AI 模型與完整系統整合 [10]；系統整合暨智慧機器人組亦以機電系統整合為主要研究方向。系內相關研究包含控制與訊號處理、機器人核心技術、感知學習與人機合作、智慧型控制、資機電整合、智慧物聯網、機器人系統與感測器等面向，與本計畫的技術鏈高度相符。我的生物醫學工程背景與穿戴式上肢輔助系統經驗，使我能由人體動作量測及實際安全需求定義問題；進入工科所後，希望進一步整合感測、控制、通訊與嵌入式韌體，建立可在實體機器人上驗證的可靠人機互動系統，並朝機器人系統與韌體工程領域發展。")

# Replace the two school-specific references with one Engineering Science source.
for p in list(doc.paragraphs):
    if p.text.startswith("[10]"):
        replace_paragraph(p, "[10] 國立成功大學工程科學系（2026）。系所特色與智慧機器人方向。https://www.es.ncku.edu.tw/esncku/zh/highschool")
    elif p.text.startswith("[11]"):
        remove_paragraph(p)

doc.core_properties.title = doc.paragraphs[0].text
doc.core_properties.subject = "成大工程科學系系統整合暨智慧機器人組研究計畫"
doc.core_properties.keywords = "智慧機器人, 機電整合, IMU, 嵌入式系統, 遙操作, 安全控制"
doc.save(OUT)
format_document(OUT)
print(OUT)
