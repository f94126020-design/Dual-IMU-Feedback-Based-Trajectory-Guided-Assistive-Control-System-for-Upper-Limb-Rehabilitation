from pathlib import Path

from docx import Document

from format_research_plans import format_document


PATH = Path(r"C:\Users\USER\Desktop\NCKU\畢業專題\成大醫工研究計畫_最終修訂.docx")


def replace(paragraph, text):
    paragraph.clear()
    paragraph.add_run(text)


doc = Document(PATH)

replace(doc.paragraphs[1], "成大醫工醫學電子組　神經肌肉訊號處理、穿戴式生醫電子與嵌入式系統")
replace(doc.paragraphs[2], "申請人：郭于綸")

replace(doc.paragraphs[4], "穿戴式生醫電子若要長時間應用於神經復健、動作輔助與居家健康照護，除了辨識準確率，還必須兼顧訊號可靠度、即時性與能源消耗。表面肌電（sEMG）是反映運動單元募集與肌肉活化的非侵入式周邊神經肌肉訊號，可在肢體產生明顯位移前提供動作意圖資訊，但容易受到電極位移、皮膚接觸、肌肉疲勞與動作偽影影響；慣性量測單元（IMU）則描述實際肢體運動，可提供互補的動作學資訊。既有研究已證明 sEMG 與慣性感測融合能改善動作辨識 [1–2]，也分別探討跨次量測變異 [3–4]、自適應取樣 [5–6] 與低功耗硬體辨識 [7–8]，但較少在同一個穿戴式系統中，將即時訊號品質、多模態融合與硬體運算負擔共同納入設計。本研究擬建立同步 IMU–sEMG 穿戴式量測平台，以可解釋的訊號品質指標動態調整融合權重與系統運作模式，再將演算法定點化並部署於低功耗 MCU 或 FPGA，建立可延伸至神經肌肉人機介面、醫療量測系統與穿戴式生醫電子的實作架構。")

replace(doc.paragraphs[7], "我目前就讀成功大學生物醫學工程學系四年級，並在生醫訊號與神經工程實驗室進行雙 IMU 上肢復健輔助系統專題，整合人體動作感測、Raspberry Pi 軟體、ESP32 韌體、Web 介面、馬達控制、通訊與安全狀態機。除了 IMU 系統，我亦進行 sEMG 動作偵測與分類、前端放大電路、PCB 製作及 SPI 資料傳輸，並實際分析訊號品質與模型泛化問題。這些經驗使我從神經工程與動作輔助需求出發，逐步建立從微弱生理訊號、類比前端、數位演算法到嵌入式系統的整合能力，也讓我理解：若忽略取樣同步、電極狀態、封包遺失、運算延遲與硬體限制，即使離線模型表現良好，也不一定能成為可靠的醫療電子系統。未來我希望深化生醫訊號演算法、嵌入式即時系統與穿戴式生醫電子三項能力，並朝韌體工程與醫療電子系統開發方向發展。")

for p in doc.paragraphs:
    if p.text.startswith("以一至兩顆 IMU 搭配雙通道或多通道 sEMG"):
        replace(p, "以一至兩顆 IMU 搭配雙通道或多通道 sEMG 為基本配置。先使用商用生醫訊號前端或既有開發板建立可重複的量測基準，完成電極保護、儀表放大、帶通／陷波、ADC、共同時間戳與封包序號，並記錄增益、飽和、電極脫落與雜訊狀態。類比前端以輸入雜訊、共模拒斥比、頻寬、增益、動態範圍及功耗作為主要規格，再依量測結果評估可程式化增益、低功耗濾波或事件偵測電路，建立由系統需求回推生醫電子前端設計的流程。")
    elif p.text.startswith("sEMG 先進行去直流、帶通／陷波"):
        replace(p, "sEMG 先進行去直流、帶通／陷波、整流及時域／頻域特徵擷取，候選品質指標包含飽和比例、基線漂移、50／60 Hz 能量比、訊號變異、通道一致性與電極接觸狀態；IMU 則檢查加速度模長、角速度突變、姿態濾波殘差、時間戳與封包完整性。品質分數先以歸一化特徵及明確門檻建立可解釋基準，再比較輕量統計分類器，並以人工加入與實際產生的異常資料驗證其區辨能力。")
    elif p.text.startswith("任務先限定為上肢靜止、屈曲、伸展"):
        replace(p, "研究先聚焦上肢靜止、屈曲、伸展與數種基本動作。基準方法包括閾值法、IMU-only、sEMG-only、固定 early／late fusion，以及 LDA、SVM 或輕量神經網路。主方法依兩種感測器的品質分數調整特徵或決策權重；若任一模態持續不可信，則輸出低信心、要求重新校正或暫停控制。以滑動視窗進行線上推論，分別計算動作起始偵測延遲、穩態分類效能及跨次穿戴穩定度。")
    elif p.text.startswith("先在 MCU 建立可量測的即時韌體基準"):
        replace(p, "在 MCU 建立可量測的即時韌體基準，將感測器驅動、共同時間戳、DMA／中斷取樣、環形緩衝區、前處理、特徵擷取、融合、分類與通訊拆分為明確模組，記錄各階段執行時間、記憶體與資料遺失。低活動狀態僅維持必要的品質監測與低成本特徵；偵測到肌肉或動作活動後，再啟動完整取樣與融合。接著比較 8、12、16 位定點格式，決定乘加器、緩衝記憶體與資料路徑規格；完成 MCU 驗證後，再以 FPGA 實作重複性高的濾波、視窗、特徵與矩陣運算，評估硬體加速效益。")
    elif p.text == "七、研究環境需求與未來學習方向":
        replace(p, "七、與成大醫工醫學電子組之契合")
    elif p.text.startswith("本計畫需要結合神經肌肉訊號、生醫訊號演算法"):
        replace(p, "成大醫工醫學電子組的研究方向涵蓋醫學電子與資訊、醫療量測與系統及生醫感測，與本計畫由 sEMG／IMU 感測、訊號品質估測、動作意圖辨識到嵌入式硬體部署的技術鏈相符。我的生醫訊號與神經工程實驗室經驗，使我已具備人體動作量測、sEMG 前端電路、PCB、SPI、嵌入式韌體與控制系統整合的基礎；進入研究所後，希望進一步學習微弱生理訊號擷取、醫療量測系統、低功耗資料路徑與硬體化生醫訊號演算法，建立三者相互連結的能力：以生醫訊號演算法理解人體動作意圖，以嵌入式系統完成可靠即時運算，再以生醫電子系統整合感測前端、運算平台與醫療應用。")

# Add an official program-direction reference without naming individual faculty.
last_ref = next(p for p in reversed(doc.paragraphs) if p.text.startswith("[8]"))
new_p = doc.add_paragraph(style=last_ref.style)
new_p._p.addprevious(last_ref._p.getnext()) if False else None
replace(new_p, "[9] 國立成功大學生物醫學工程學系（2026）。碩士班醫學電子組研究方向：醫學電子與資訊、醫療量測與系統、生醫感測。")

doc.core_properties.title = doc.paragraphs[0].text
doc.core_properties.subject = "成大醫工醫學電子組研究計畫"
doc.core_properties.keywords = "sEMG, IMU, 生醫訊號處理, 穿戴式生醫電子, 嵌入式系統, 醫療量測"
doc.save(PATH)
format_document(PATH)
print(PATH)
