from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"C:\Users\USER\Desktop\NCKU\畢業專題")
ROOT.mkdir(parents=True, exist_ok=True)

FONT = "Microsoft JhengHei"

def set_font(run, size=11, bold=False):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def borders(table):
    tblPr = table._tbl.tblPr
    node = tblPr.find(qn("w:tblBorders"))
    if node is None:
        node = OxmlElement("w:tblBorders")
        tblPr.append(node)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "D9D9D9")
        node.append(el)

def base_doc(title, subtitle):
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
    sec.left_margin, sec.right_margin = Cm(2.3), Cm(2.3)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(6)
    for name, size in [("Title", 20), ("Heading 1", 15), ("Heading 2", 12)]:
        s = styles[name]
        s.font.name = FONT
        s._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
        s.paragraph_format.space_after = Pt(6)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(title), 20, True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(subtitle), 11)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("申請人：郭于綸　｜　工作稿：文獻探勘與研究架構版"), 10)
    return doc

def h(doc, text, level=1):
    doc.add_paragraph(text, style=f"Heading {level}")

def para(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        set_font(p.add_run(bold_lead), 10.5, True)
        set_font(p.add_run(text[len(bold_lead):]), 10.5)
    else:
        set_font(p.add_run(text), 10.5)
    return p

def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        set_font(p.add_run(item), 10.5)

def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    borders(t)
    for i, text in enumerate(headers):
        c=t.rows[0].cells[i]; shade(c,"D9EAF7"); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_font(p.add_run(text),9,True)
    for ridx,row in enumerate(rows):
        cells=t.add_row().cells
        if ridx%2: 
            for c in cells: shade(c,"F7FAFC")
        for i,text in enumerate(row):
            cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cells[i].paragraphs[0]; set_font(p.add_run(str(text)),9)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Cm(w)
    doc.add_paragraph()
    return t

def add_refs(doc, refs):
    h(doc, "參考文獻", 1)
    for i,ref in enumerate(refs,1):
        para(doc, f"[{i}] {ref}")

def build_a():
    doc=base_doc("基於雙 IMU 狀態估測之繩索驅動上肢輔助機器人自適應按需控制", "A 版　機器人控制、訊號處理與嵌入式即時實現")
    h(doc, "摘要")
    para(doc, "繩索驅動上肢輔助機器人具有低慣量、驅動器可遠端配置與機構彈性高等優點，但繩索鬆弛、摩擦、遲滯、人體參數差異及感測延遲會降低狀態估測與控制可靠度。本研究擬以雙 IMU 及嵌入式控制平台為基礎，先建立可量測的關節角度與致動狀態估測，再以擾動觀測與自適應輔助機制，在使用者無法維持目標動作時提供有限輔助。研究將以 PID 與固定門檻方法為基準，比較不同感測、負載、摩擦及延遲條件下的估測誤差、軌跡追蹤、輔助介入與即時運算表現，並先以固定架與假肢平台驗證，不預設臨床效果。")
    h(doc, "一、研究背景與問題")
    para(doc, "現有文獻已將繩索驅動復健機器人的控制分為 PID、阻抗／導納、按需輔助與自適應等方法。然而，高階控制表現依賴可信的關節狀態與人機互動估測。低成本 IMU 受感測器安裝、段體校正、軟組織移動、零偏與磁干擾影響；繩索傳動又存在單向拉力、摩擦與鬆弛。若僅依賴 PWM 命令或角度誤差，將難以判斷使用者主動貢獻與實際輔助狀態。")
    para(doc, "文獻顯示，解剖約束、功能性校正與卡爾曼濾波可改善上肢關節角度估測；非線性擾動觀測器可降低繩索機器人受未建模擾動影響；按需輔助則強調只在必要時介入。但三者常被分開處理，且對低成本嵌入式平台的取樣同步、計算延遲與安全狀態缺乏整合評估。")
    h(doc, "二、研究目的")
    bullets(doc, [
        "建立雙 IMU 的功能性校正、同步取樣與關節角度／角速度估測流程。",
        "建立繩索驅動系統簡化動態模型，辨識摩擦、遲滯、鬆繩與負載變化。",
        "設計狀態／擾動觀測器，比較其相對於現行濾波與 PWM-time 估算的改善。",
        "將估測結果用於自適應按需輔助，降低不必要介入並維持追蹤與安全。",
        "在 MCU 與嵌入式 Linux 架構上實現，量化運算延遲、控制迴圈抖動及資源使用。"
    ])
    h(doc, "三、研究問題與假設")
    table(doc,["研究問題","可檢驗假設","主要指標"],[
        ("RQ1：校正與解剖約束是否改善雙 IMU 估測？","相較單次靜態歸零，功能性校正與狀態估測可降低不同戴法及動作速度下的誤差。","MAE、RMSE、drift、延遲"),
        ("RQ2：觀測器能否補償未建模的繩索與負載擾動？","加入擾動估測後，不同負載與摩擦下的追蹤誤差與輸出抖動將降低。","RMSE、峰值誤差、恢復時間、輸出總變差"),
        ("RQ3：自適應輔助是否比固定門檻更符合按需原則？","在追蹤誤差不顯著惡化下，降低輔助時間與誤觸發。","AAN 介入比例、誤／漏觸發、追蹤誤差"),
        ("RQ4：嵌入式延遲如何影響控制效果？","時間戳、固定週期排程與安全節點可降低 jitter 對估測與控制的影響。","WCET、jitter、封包遺失、安全停止時間")
    ],[4.8,7.3,4.0])
    h(doc, "四、研究方法")
    h(doc, "4.1 平台與量測",2)
    para(doc, "以現有上臂／前臂雙 IMU、Raspberry Pi、ESP32 與三路馬達平台為起點。新增外部角度參考、捲線輪編碼器、馬達電流與代表性張力量測，並統一時間戳、設定取樣週期及儲存原始資料。")
    h(doc, "4.2 IMU 校正與狀態估測",2)
    para(doc, "比較靜態零點、N-pose 與功能性關節軸校正。基準為互補濾波，改良方法採 EKF 或 UKF，將關節角度、角速度與陀螺儀零偏納入狀態，並引入肘關節解剖約束。以外部角度或視覺追蹤為參考，評估戴法偏差、動作速度與馬達干擾。")
    h(doc, "4.3 系統辨識與擾動觀測",2)
    para(doc, "以階躍、正弦、Chirp 與 PRBS 輸入取得不同方向、負載與預張力下的資料，建立可解釋的低階模型。採擾動觀測器估測摩擦、重力、人體負載與未建模影響；若張力量測充足，再評估電流／編碼器的間接張力估測。")
    h(doc, "4.4 控制器與按需輔助",2)
    para(doc, "以含 anti-windup、導數濾波與飽和處理的 PID 為基準。主方法為「擾動觀測器輔助控制＋自適應 AAN 決策」：觀測器補償未建模擾動，AAN 依持續誤差、動作趨勢與估測擾動調整門檻或輔助增益。不將 PID、LADRC、ILC、MPC 與強化學習全部列為主方法；ILC 僅在重複軌跡研究需要時作延伸。")
    h(doc, "4.5 嵌入式即時實現與安全",2)
    para(doc, "ESP32 執行固定週期的輸出更新、方向映射、斜率限制、watchdog、timeout 與急停；Raspberry Pi 執行估測、控制與資料記錄。記錄感測時間戳、控制開始／結束、通訊往返與馬達更新時刻，以分析延遲鏈及最壞執行時間。")
    h(doc, "五、實驗設計")
    table(doc,["階段","測試條件","比較方法","指標"],[
        ("感測驗證","多種戴法、靜態／動態、不同速度","CF vs. EKF/UKF；靜態 vs. 功能校正","MAE、RMSE、drift、延遲"),
        ("致動台架","負載、摩擦、鬆繩、正反轉","PWM-time vs. 編碼器／電流／張力估測","估測誤差、遲滯、鬆繩事件"),
        ("固定架控制","階躍、正弦、重複軌跡、外加擾動","PID vs. PID+觀測器","RMSE、穩定時間、輸出抖動"),
        ("按需輔助","模擬不同主動能力與負載","無輔助、固定門檻、自適應 AAN","介入比例、誤／漏觸發、追蹤誤差"),
        ("即時與故障","加入延遲、jitter、斷訊與感測異常","現行架構 vs. 時間戳／固定週期架構","WCET、jitter、停止時間、超限事件")
    ],[2.2,4.8,5.0,4.1])
    h(doc, "六、預期貢獻")
    bullets(doc,[
        "一套可重複的雙 IMU 校正、時間同步與上肢關節狀態估測流程。",
        "一個將繩索傳動擾動估測與按需輔助連結的可解釋控制架構。",
        "感測誤差、控制表現、介入程度與嵌入式延遲的量化證據。",
        "可追溯的實驗資料、程式版本、參數與故障紀錄。"
    ])
    h(doc, "七、風險與替代方案")
    table(doc,["風險","控制方式／替代方案"],[
        ("外部真值系統不可得","先以編碼器、電子量角器或視覺標記建立單平面參考。"),
        ("張力感測器數量不足","以一通道作校正參考，研究電流／編碼器間接估測，不宣稱等同真實輔助力。"),
        ("人體模型過於複雜","先限定單自由度與假肢／固定架，使用集總擾動表示未建模效應。"),
        ("進階控制器無法穩定實現","保留 PID+擾動補償為最小可行方法，先完成實驗可重複性。"),
        ("無法進行人體試驗","以台架、假肢與擾動注入完成工程驗證；人體測試依倫理與實驗室條件另定。")
    ],[5.0,11.1])
    h(doc, "八、進度規畫")
    table(doc,["期程","主要工作"],[
        ("碩一上","文獻整理、數學與控制基礎、量測平台、IMU校正與外部真值。"),
        ("碩一下","系統辨識、觀測器設計、基準控制實驗。"),
        ("碩二上","自適應 AAN、嵌入式實現、完整對照與故障測試。"),
        ("碩二下","補充實驗、統計分析、論文與口試；視條件規畫健康受試者驗證。")
    ],[3.2,12.9])
    refs=[
        "Shoaib, M., et al. (2021). Cable Driven Rehabilitation Robots: Comparison of Applications and Control Strategies. IEEE Access, 9, 110396–110420. https://doi.org/10.1109/ACCESS.2021.3102107",
        "Proietti, T., et al. (2021). Review on Patient-Cooperative Control Strategies for Upper-Limb Rehabilitation Exoskeletons. Frontiers in Robotics and AI, 8, 745018. https://doi.org/10.3389/frobt.2021.745018",
        "Delgado, P., & Yihun, Y. (2023). Integration of Task-Based Exoskeleton with an Assist-as-Needed Algorithm for Patient-Centered Elbow Rehabilitation. Sensors, 23(5), 2460. https://doi.org/10.3390/s23052460",
        "Wu, Q., et al. (2018). Patient-Active Control of a Powered Exoskeleton Targeting Upper Limb Rehabilitation Training. Frontiers in Neurology, 9, 817. https://doi.org/10.3389/fneur.2018.00817",
        "Liu, J., et al. (2017). Sliding Mode Tracking Control of a Wire-Driven Upper-Limb Rehabilitation Robot with Nonlinear Disturbance Observer. Frontiers in Neurology, 8, 646. https://doi.org/10.3389/fneur.2017.00646",
        "Passon, A., et al. (2020). Inertial-Robotic Motion Tracking in End-Effector-Based Rehabilitation Robots. Frontiers in Robotics and AI, 7, 554639. https://doi.org/10.3389/frobt.2020.554639",
        "Bicchi, A., et al. (2024). Improved Estimation of Elbow Flexion Angle from IMU Measurements Using Anatomical Constraints. IRBM, 45(1), 100820. https://doi.org/10.1016/j.irbm.2024.100820",
        "Poitras, I., et al. (2024). Effects of IMU sensor-to-segment calibration on clinical 3D elbow joint angles estimation. Journal of Biomechanics. PMID: 38835976.",
        "Influence of Visual-Inertial Sensor-to-Segment Calibration on Upper Limb Joint Angles Estimation From Multiple Inverse Kinematics Methods (2025). IEEE T-ASE, 22, 11519–11528. https://doi.org/10.1109/TASE.2025.3535857",
        "Seyfi, N. S., & Khalaji, A. K. (2022). Robust control of a cable-driven rehabilitation robot for lower and upper limbs. ISA Transactions, 125, 268–289. https://doi.org/10.1016/j.isatra.2021.07.016",
        "Wang, Y., et al. (2021). Control strategy and experimental research of a cable-driven lower limb rehabilitation robot. Proceedings of the Institution of Mechanical Engineers, Part C. https://doi.org/10.1177/0954406220952510",
        "Reference-model-based control including human torque estimation for cable-driven rehabilitation system (2026). IFAC Journal of Systems and Control, 35, 100403. https://doi.org/10.1016/j.ifacsc.2026.100403"
    ]
    add_refs(doc,refs)
    doc.save(ROOT/"A版_機器人控制_IMU按需輔助_研究計畫工作稿.docx")

def build_b():
    doc=base_doc("基於 IMU 與表面肌電融合之上肢動作意圖辨識與低功耗嵌入式實現", "B 版　生醫訊號處理、穿戴電子與可硬體化架構")
    h(doc, "摘要")
    para(doc, "單一 IMU 可直接反映肢體運動，但難以在動作尚未明顯前估測肌肉啟動；表面肌電可提供較早的肌肉活化資訊，卻受電極位置、接觸、動作偽影與個體差異影響。本研究擬建立 IMU與sEMG同步取樣系統，比較單一感測與多模態融合對上肢動作起始、方向與狀態辨識的影響，再將方法定點化與輕量化，部署於低功耗 MCU 或 FPGA 平台。研究將同時評估辨識效能、延遲、記憶體、運算量與功耗，以建立可向低功耗生醫訊號處理電路延伸的系統架構。")
    h(doc, "一、研究動機與定位")
    para(doc, "本計畫以「訊號處理演算法的低功耗嵌入式與可硬體化實現」為主，不預設申請人已具備完整類比 IC 或 ASIC 流片經驗。對偏生醫訊號的教授，強化訊號品質、個人化與模型驗證；對偏 IC 與穿戴電子的教授，強化類比前端、ADC、定點數字濾波、特徵擷取加速與功耗分析。")
    h(doc, "二、研究問題")
    bullets(doc,[
        "IMU 與 sEMG 融合是否能在不顯著增加延遲下，改善動作起始與方向辨識？",
        "如何處理電極位置、跨日漂移、動作偽影與不同使用者所造成的分佈差異？",
        "浮點模型轉換為定點與輕量化架構後，準確率、延遲、記憶體與功耗間有何權衡？"
    ])
    h(doc, "三、研究方法")
    table(doc,["模組","預定內容","比較基準"],[
        ("訊號擷取","雙 IMU、雙通道或多通道 sEMG、共同時間戳與標記。","IMU-only、sEMG-only、fusion"),
        ("前處理","sEMG 帶通／陷波、包絡、視窗；IMU 姿態與動作特徵；訊號品質偵測。","固定參數 vs. 個人化正規化"),
        ("辨識","動作起始、屈／伸方向與狀態；LDA/SVM 或輕量網路。","門檻法 vs. 傳統機器學習 vs. 輕量模型"),
        ("嵌入式與硬體化","定點化、記憶體配置、DMA／buffer、軟硬體分工，視條件使用 FPGA。","浮點 PC vs. MCU/FPGA"),
        ("評估","跨受試者、跨日、不同電極／感測器位置與動作速度。","F1、誤觸發、偵測延遲、RAM/Flash、能耗")
    ],[3.0,8.0,5.1])
    h(doc, "四、預期貢獻")
    bullets(doc,[
        "一套 IMU／sEMG 同步與訊號品質管理流程。",
        "多模態融合在辨識準確度、提前量與可靠性上的量化比較。",
        "具明確運算、記憶體與功耗預算的嵌入式實現。",
        "從演算法至定點數字處理模組／硬體加速器的可延伸設計依據。"
    ])
    h(doc, "五、進度與可行性")
    table(doc,["期程","主要工作"],[
        ("碩一上","文獻、擷取硬體、同步與訊號品質驗證。"),
        ("碩一下","資料收集、特徵與融合模型、跨受試者評估。"),
        ("碩二上","定點化、MCU/FPGA 部署、延遲與功耗測試。"),
        ("碩二下","完整對照、模組化架構、論文與成果整理。")
    ],[3.2,12.9])
    h(doc, "六、待完成的專屬文獻探勘")
    para(doc, "本版先建立可行的研究骨架。後續應依申請教授分成「生醫訊號與意圖辨識」及「低功耗混合訊號／訊號處理 IC」兩條文獻線，再決定主要貢獻是模型驗證、類比前端、數字處理或硬體加速。")
    refs=[
        "Li, H., et al. (2025). A Home-based Dual-mode Upper Limb Rehabilitation System: Teleoperation Mode and Bilateral Mode with sEMG and IMU. IEEE Journal of Biomedical and Health Informatics, 29(11), 8140–8152. https://doi.org/10.1109/JBHI.2025.3588404",
        "Delgado, P., & Yihun, Y. (2023). Integration of Task-Based Exoskeleton with an Assist-as-Needed Algorithm for Patient-Centered Elbow Rehabilitation. Sensors, 23(5), 2460. https://doi.org/10.3390/s23052460",
        "Assist-As-Needed Exoskeleton for Hand Joint Rehabilitation Based on Muscle Effort Detection (2021). Sensors. PMID: 34206714.",
        "A Multi-Modal Under-Sensorized Wearable System for Optimal Kinematic and Muscular Tracking of Human Upper Limb Motion (2023). Sensors. PMCID: PMC10098930."
    ]
    add_refs(doc,refs)
    doc.save(ROOT/"B版_生醫訊號_低功耗嵌入式_研究計畫工作稿.docx")

if __name__ == "__main__":
    build_a()
    build_b()
    print("done")
