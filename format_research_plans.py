from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BODY_CJK = "PMingLiU"
HEAD_CJK = "Microsoft JhengHei"
LATIN = "Times New Roman"


def set_run_font(run, size=None, bold=None, cjk=BODY_CJK, latin=LATIN):
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), latin)
    fonts.set(qn("w:hAnsi"), latin)
    fonts.set(qn("w:eastAsia"), cjk)
    fonts.set(qn("w:cs"), latin)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def set_style_font(style, size, bold=False, cjk=BODY_CJK):
    style.font.name = LATIN
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), LATIN)
    fonts.set(qn("w:hAnsi"), LATIN)
    fonts.set(qn("w:eastAsia"), cjk)
    fonts.set(qn("w:cs"), LATIN)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tcpr = tc.get_or_add_tcPr()
    mar = tcpr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tcpr.append(mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tblpr = table._tbl.tblPr
    old = tblpr.find(qn("w:tblBorders"))
    if old is not None:
        tblpr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "5")
        el.set(qn("w:color"), "B7B7B7")
        borders.append(el)
    tblpr.append(borders)


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def prevent_row_split(row):
    trpr = row._tr.get_or_add_trPr()
    if trpr.find(qn("w:cantSplit")) is None:
        trpr.append(OxmlElement("w:cantSplit"))


def add_page_number(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.clear()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    run = p.add_run()
    set_run_font(run, 9, cjk=HEAD_CJK)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    for el in (fld_begin, instr, fld_sep, text, fld_end):
        run._r.append(el)


def format_document(path):
    doc = Document(path)

    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.4)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.7)
        section.right_margin = Cm(2.7)
        section.header_distance = Cm(1.2)
        section.footer_distance = Cm(1.2)
        add_page_number(section)

    styles = doc.styles
    normal = styles["Normal"]
    set_style_font(normal, 12, False, BODY_CJK)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.first_line_indent = Pt(24)
    normal.paragraph_format.widow_control = True

    title_style = styles["Title"]
    set_style_font(title_style, 18, True, HEAD_CJK)
    title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_style.paragraph_format.space_before = Pt(0)
    title_style.paragraph_format.space_after = Pt(8)
    title_style.paragraph_format.keep_with_next = True

    h1 = styles["Heading 1"]
    set_style_font(h1, 14, True, HEAD_CJK)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)
    h1.paragraph_format.keep_with_next = True
    h1.paragraph_format.keep_together = True
    h1.paragraph_format.first_line_indent = Pt(0)

    h2 = styles["Heading 2"]
    set_style_font(h2, 12, True, HEAD_CJK)
    h2.paragraph_format.space_before = Pt(9)
    h2.paragraph_format.space_after = Pt(4)
    h2.paragraph_format.keep_with_next = True
    h2.paragraph_format.keep_together = True
    h2.paragraph_format.first_line_indent = Pt(0)

    for list_name in ("List Bullet", "List Number"):
        if list_name in styles:
            s = styles[list_name]
            set_style_font(s, 12, False, BODY_CJK)
            s.paragraph_format.left_indent = Cm(0.8)
            s.paragraph_format.first_line_indent = Cm(-0.4)
            s.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            s.paragraph_format.space_after = Pt(2)

    heading_seen = False
    reference_mode = False
    for index, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if index == 0:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.keep_with_next = True
            for run in p.runs:
                set_run_font(run, 18, True, HEAD_CJK)
            continue
        if text == "參考文獻":
            reference_mode = True
        if p.style.name.startswith("Heading"):
            heading_seen = True
            p.paragraph_format.first_line_indent = Pt(0)
            for run in p.runs:
                set_run_font(run, 14 if p.style.name == "Heading 1" else 12, True, HEAD_CJK)
            continue
        if not heading_seen:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.line_spacing = 1.2
            p.paragraph_format.space_after = Pt(3)
            size = 11 if index == 1 else 10.5
            for run in p.runs:
                set_run_font(run, size, False, HEAD_CJK)
            continue
        if reference_mode and text.startswith("["):
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(-0.75)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(5)
            for run in p.runs:
                set_run_font(run, 10, False, BODY_CJK)
            continue
        if p.style.name not in ("List Bullet", "List Number"):
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Pt(24)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            p.paragraph_format.space_after = Pt(4)
        for run in p.runs:
            set_run_font(run, 12, run.bold, BODY_CJK)

    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        set_table_borders(table)
        for ridx, row in enumerate(table.rows):
            prevent_row_split(row)
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margins(cell)
                shade(cell, "D9E2F3" if ridx == 0 else ("F7F7F7" if ridx % 2 == 0 else "FFFFFF"))
                for p in cell.paragraphs:
                    p.paragraph_format.first_line_indent = Pt(0)
                    p.paragraph_format.line_spacing = 1.15
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(1)
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ridx == 0 else WD_ALIGN_PARAGRAPH.LEFT
                    for run in p.runs:
                        set_run_font(run, 10, ridx == 0, HEAD_CJK if ridx == 0 else BODY_CJK)

    props = doc.core_properties
    props.title = doc.paragraphs[0].text.strip()
    props.subject = "研究計畫"
    props.keywords = "研究計畫, 機器人控制, 生醫訊號, 嵌入式系統"

    doc.save(path)


if __name__ == "__main__":
    root = Path(r"C:\Users\USER\Desktop\NCKU\畢業專題\正式排版研究計畫")
    for docx in sorted(root.glob("*.docx")):
        format_document(docx)
        print(docx)
