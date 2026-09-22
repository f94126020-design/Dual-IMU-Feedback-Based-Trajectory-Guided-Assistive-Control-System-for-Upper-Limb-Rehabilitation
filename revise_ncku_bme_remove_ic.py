from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


SOURCE = Path(r"C:\Users\USER\Desktop\NCKU\推甄準備\成大醫工醫學電子組\成大醫工研究計畫_最終版.docx")
OUTPUT = Path(r"C:\Users\USER\Desktop\NCKU\畢業專題\成大醫工研究計畫_最終版.docx")


def set_text_preserving_paragraph(paragraph, text):
    paragraph.clear()
    paragraph.add_run(text)


def set_bilingual_font(run):
    run.font.name = "Times New Roman"
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:cs"), "Times New Roman")
    rfonts.set(qn("w:eastAsia"), "標楷體")


doc = Document(SOURCE)

# Remove the chip-design deliverable and keep the proposal at the biomedical
# electronics / embedded-system level that can be executed with available resources.
set_text_preserving_paragraph(
    doc.paragraphs[43],
    "形成可供後續開發低雜訊生醫訊號擷取前端、可程式化資料轉換、嵌入式生醫訊號處理模組，或神經肌肉人機介面之系統規格與設計依據。",
)

# State the cross-college learning plan without treating IC design as a promised
# research output or as an existing undergraduate capability.
fit = doc.paragraphs[46]
old = fit.text
addition = (
    "大學階段的訓練以生醫工程、訊號處理與系統整合為主，尚未接受完整的晶片設計訓練；"
    "若研究後續需延伸至積體電路或更底層的硬體實現，我將主動跨修電資學院相關課程，補足數位系統、硬體描述語言與積體電路設計基礎。"
    "本研究的可執行核心仍以商用生醫訊號前端、MCU 或 FPGA 平台完成演算法驗證與嵌入式系統整合。"
)
set_text_preserving_paragraph(fit, old + addition)

# Replace the former IC-schedule risk with a realistic hardware-integration risk.
doc.tables[4].cell(5, 0).text = "硬體化開發時程不足"
doc.tables[4].cell(5, 1).text = "以 MCU 定點實現作為最小可行成果；FPGA 加速則依研究進度與平台資源逐步深化。"

# Remove IC basics from the first-semester commitment and move it to an optional,
# explicit cross-college course plan.
doc.tables[5].cell(1, 1).text = "補強生醫電子、數位訊號處理與嵌入式系統基礎；完成同步擷取、前端規格及基準資料。若研究需求涉及積體電路層級，另跨修電資學院相關課程。"

# Apply the requested bilingual font policy throughout body text, tables,
# headers, footers, and styles while preserving existing sizes and hierarchy.
for style in doc.styles:
    if hasattr(style, "font"):
        style.font.name = "Times New Roman"
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "標楷體")
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Times New Roman")
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Times New Roman")

for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        set_bilingual_font(run)

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_bilingual_font(run)

for section in doc.sections:
    for part in (section.header, section.footer):
        for paragraph in part.paragraphs:
            for run in paragraph.runs:
                set_bilingual_font(run)
        for table in part.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            set_bilingual_font(run)

doc.core_properties.title = "穿戴式 IMU 與表面肌電多模態動作意圖辨識及低功耗嵌入式架構"
doc.core_properties.subject = "成大醫工醫學電子組研究計畫"
doc.save(OUTPUT)

print(OUTPUT)
